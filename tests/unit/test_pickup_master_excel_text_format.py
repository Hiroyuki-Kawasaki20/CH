"""入車時間マスタの文字列セル・入力用空行・固定名テーブルの保存回帰。"""
from datetime import datetime, time

import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from src.services import data_loader
from src.services.data_loader import (
    MasterFileError,
    load_pickup_time_master_xlsx,
    save_pickup_time_master_xlsx,
)


HEADERS = ["OData_納入先", "NONYUHIBIN", "入車時間", "セットありフラグ"]
TEXT_COLUMNS = ["NONYUHIBIN", "入車時間"]
SHEET_NAME = "入車時間マスタ"
TABLE_NAME = "PickupTimeMaster"


def _master_df():
    return pd.DataFrame([
        ["テスト納入先", "01", "08:30", "1"],
        ["テスト納入先", "02", "25:10", "0"],
        ["テスト納入先", "10", "09:15", ""],
    ], columns=HEADERS)


def _text_cells(sheet, row):
    positions = {cell.value: cell.column for cell in sheet[1]}
    return [sheet.cell(row, positions[name]) for name in TEXT_COLUMNS]


def test_values_remain_string_cells(tmp_path):
    source = _master_df()
    path = tmp_path / "master.xlsx"
    save_pickup_time_master_xlsx(source, path)
    book = load_workbook(path)
    try:
        assert book.sheetnames == [SHEET_NAME]
        sheet = book[SHEET_NAME]
        assert sheet.max_column == len(HEADERS)  # pandasのindex列は出力しない
        for row_index, expected in enumerate(
            [("01", "08:30"), ("02", "25:10"), ("10", "09:15")], start=2
        ):
            cells = _text_cells(sheet, row_index)
            assert tuple(cell.value for cell in cells) == expected
            assert all(isinstance(cell.value, str) and cell.data_type == "s" for cell in cells)
    finally:
        book.close()


@pytest.mark.parametrize("column", TEXT_COLUMNS)
def test_existing_data_cells_have_text_format(tmp_path, column):
    path = tmp_path / "master.xlsx"
    save_pickup_time_master_xlsx(_master_df(), path)
    book = load_workbook(path)
    try:
        sheet = book[SHEET_NAME]
        column_index = HEADERS.index(column) + 1
        assert [sheet.cell(row, column_index).number_format for row in range(2, 5)] == ["@"] * 3
    finally:
        book.close()


def test_second_save_preserves_values_and_formats_added_row(tmp_path):
    path = tmp_path / "master.xlsx"
    source = _master_df()
    save_pickup_time_master_xlsx(source.iloc[:1], path)
    # GUI行追加に相当する文字列行を追加し、同じファイルを再生成する。
    save_pickup_time_master_xlsx(source.iloc[:2], path)
    book = load_workbook(path)
    try:
        sheet = book[SHEET_NAME]
        for row, expected in [(2, ("01", "08:30")), (3, ("02", "25:10"))]:
            cells = _text_cells(sheet, row)
            assert tuple(cell.value for cell in cells) == expected
            assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells)
    finally:
        book.close()


@pytest.mark.parametrize("row_count", [0, 2])
def test_exactly_100_blank_input_rows_have_only_text_format(tmp_path, row_count):
    path = tmp_path / "master.xlsx"
    save_pickup_time_master_xlsx(_master_df().iloc[:row_count], path)
    book = load_workbook(path)
    try:
        sheet = book[SHEET_NAME]
        last_data_row = row_count + 1
        assert sheet.max_row == last_data_row + 100
        for row in range(last_data_row + 1, last_data_row + 101):
            assert all(sheet.cell(row, col).value is None for col in range(1, 5))
            assert all(cell.number_format == "@" for cell in _text_cells(sheet, row))
            assert sheet.cell(row, 1).number_format == "General"
            assert sheet.cell(row, 4).number_format == "General"
        assert all(cell.number_format == "General" for cell in _text_cells(sheet, last_data_row + 101))
    finally:
        book.close()


@pytest.mark.parametrize("row_count", [0, 2, 3])
def test_official_reload_excludes_formatted_blank_rows_and_preserves_other_columns(tmp_path, row_count):
    source = _master_df().iloc[:row_count].copy()
    legacy_path = tmp_path / "legacy.xlsx"
    path = tmp_path / "master.xlsx"
    # 変更前のwriterと同じ出力を比較基準にする（空フラグの既存読込結果も保持）。
    source.to_excel(legacy_path, index=False, engine="openpyxl", sheet_name=SHEET_NAME)
    expected = load_pickup_time_master_xlsx(legacy_path)
    save_pickup_time_master_xlsx(source, path)
    actual = load_pickup_time_master_xlsx(path)
    assert len(actual) == row_count
    assert list(actual.columns) == HEADERS
    assert_frame_equal(actual, expected)


@pytest.mark.parametrize("missing", [None, float("nan"), pd.NA, pd.NaT], ids=["None", "NaN", "NA", "NaT"])
def test_missing_values_remain_empty_not_literal_strings(tmp_path, missing):
    source = pd.DataFrame([["テスト納入先", missing, missing, "0"]], columns=HEADERS)
    path = tmp_path / "master.xlsx"
    save_pickup_time_master_xlsx(source, path)
    book = load_workbook(path)
    try:
        assert all(cell.value is None for cell in _text_cells(book[SHEET_NAME], 2))
    finally:
        book.close()
    loaded = load_pickup_time_master_xlsx(path)
    assert loaded.loc[0, "NONYUHIBIN"] == ""
    assert loaded.loc[0, "入車時間"] == ""
    assert loaded.loc[0, "セットありフラグ"] == "0"


def test_save_does_not_mutate_caller_dataframe(tmp_path):
    source = _master_df().iloc[::-1].copy()
    source.loc[source.index[0], "入車時間"] = None
    source["対象外の列"] = "保持"
    before = source.copy(deep=True)
    save_pickup_time_master_xlsx(source, tmp_path / "master.xlsx")
    assert_frame_equal(source, before)


@pytest.mark.parametrize("bin_value,time_value", [
    (1, time(8, 30)),
    (1.0, datetime(2026, 10, 23, 8, 30)),
    (2, 0.5),
])
def test_existing_non_string_inputs_keep_official_loader_results(tmp_path, bin_value, time_value):
    # 本番GUIは文字列を渡す。保存APIで従来受け付けた型も勝手に時刻解釈しない。
    source = pd.DataFrame([["テスト納入先", bin_value, time_value, "0"]], columns=HEADERS)
    legacy_path = tmp_path / "legacy.xlsx"
    path = tmp_path / "master.xlsx"
    source.to_excel(legacy_path, index=False, engine="openpyxl", sheet_name=SHEET_NAME)
    save_pickup_time_master_xlsx(source, path)
    assert_frame_equal(load_pickup_time_master_xlsx(path), load_pickup_time_master_xlsx(legacy_path))


@pytest.mark.parametrize("row_count", [0, 1, 2, 78])
def test_table_name_headers_range_and_filter(tmp_path, row_count):
    source = pd.DataFrame([
        ["テスト納入先", f"{index + 1:02d}", "08:30", "1"]
        for index in range(row_count)
    ], columns=HEADERS)
    path = tmp_path / "master.xlsx"
    save_pickup_time_master_xlsx(source, path)
    book = load_workbook(path)
    try:
        sheet = book[SHEET_NAME]
        assert list(sheet.tables) == [TABLE_NAME]
        table = sheet.tables[TABLE_NAME]
        assert table.ref == f"A1:D{row_count + 1}"  # 入力用100行は含めない
        assert [column.name for column in table.tableColumns] == HEADERS
        assert [cell.value for cell in sheet[1]] == HEADERS
        assert table.autoFilter is not None
        assert table.autoFilter.ref == table.ref
        assert table.tableStyleInfo is not None
    finally:
        book.close()
    loaded = load_pickup_time_master_xlsx(path)
    assert len(loaded) == row_count
    if row_count:
        assert_frame_equal(loaded, source)
    else:
        assert list(loaded.columns) == HEADERS


def test_repeated_save_has_one_table_and_tracks_growing_shrinking_data(tmp_path):
    path = tmp_path / "master.xlsx"
    for row_count in (2, 2, 3, 1, 0):
        source = _master_df().iloc[:row_count]
        save_pickup_time_master_xlsx(source, path)
        book = load_workbook(path)
        try:
            sheet = book[SHEET_NAME]
            assert list(sheet.tables) == [TABLE_NAME]
            assert sheet.tables[TABLE_NAME].ref == f"A1:D{row_count + 1}"
            assert sheet.max_row == row_count + 101
        finally:
            book.close()
        assert len(load_pickup_time_master_xlsx(path)) == row_count


def test_writer_locates_text_columns_by_name_without_mutating_input(tmp_path):
    source = _master_df()[["入車時間", "OData_納入先", "セットありフラグ", "NONYUHIBIN"]]
    before = source.copy(deep=True)
    path = tmp_path / "master.xlsx"
    data_loader._write_master_excel(source, path)
    assert_frame_equal(source, before)
    book = load_workbook(path)
    try:
        sheet = book[SHEET_NAME]
        assert all(cell.number_format == "@" for cell in _text_cells(sheet, 2))
        assert sheet["B2"].number_format == "General"
        assert sheet["C2"].number_format == "General"
    finally:
        book.close()


def test_rescue_file_has_same_formats_and_table(tmp_path, monkeypatch):
    rescue_dir = tmp_path / "rescue"
    monkeypatch.setattr(data_loader, "_master_rescue_dir", lambda: rescue_dir)
    path = tmp_path / "missing" / "master.xlsx"
    with pytest.raises(MasterFileError):
        save_pickup_time_master_xlsx(_master_df(), path)
    assert not path.parent.exists()
    rescue_files = list(rescue_dir.glob("*.xlsx"))
    assert len(rescue_files) == 1
    book = load_workbook(rescue_files[0])
    try:
        sheet = book[SHEET_NAME]
        assert sheet.tables[TABLE_NAME].ref == "A1:D4"
        assert all(cell.number_format == "@" for cell in _text_cells(sheet, 2))
    finally:
        book.close()
