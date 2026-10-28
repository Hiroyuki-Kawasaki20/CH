"""利用者別ローカル出力。import時定数は子プロセスで検証し、親を汚染しない。"""

import json
import os
from pathlib import Path
import subprocess
import sys
import textwrap

import pytest


REPO_ROOT = Path(__file__).resolve().parents[2]
_CHILD_SETUP = """
import json
import os
from pathlib import Path
import sys
from unittest.mock import patch

sandbox = Path(os.environ['CH_TEST_SANDBOX']).resolve()
expected = Path(os.environ['CH_TEST_EXPECTED'])
null_device = Path(os.devnull).resolve()
# REDでも開発者フォルダや本番データへ書き込ませない。
def guard(event, args):
    paths = []
    if event == 'open':
        path, mode, flags = args
        if isinstance(path, (str, bytes, os.PathLike)) and (
            (mode and any(c in mode for c in 'wax+'))
            or flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC)
        ):
            paths = [path]
    elif event in ('os.mkdir', 'os.remove', 'os.rmdir'):
        paths = [args[0]]
    elif event in ('os.rename', 'os.link', 'os.symlink'):
        paths = args[:2]
    for path in paths:
        resolved = Path(os.fsdecode(path)).resolve()
        if event == 'open' and resolved == null_device:
            continue
        assert resolved.is_relative_to(sandbox), (
            f'non-sandbox mutation blocked: {event} {path}'
        )
sys.addaudithook(guard)
"""


def _run_child(tmp_path, code, *, localappdata="local", trace=False):
    home = tmp_path / "home"
    home.mkdir(exist_ok=True)
    env = os.environ.copy()
    env.update({
        "USERPROFILE": str(home), "HOME": str(home),
        "TEMP": str(tmp_path), "TMP": str(tmp_path), "TMPDIR": str(tmp_path),
        "PYTHONPATH": str(REPO_ROOT), "PYTHONDONTWRITEBYTECODE": "1",
        "PYTHONIOENCODING": "utf-8", "CH_TEST_SANDBOX": str(tmp_path),
        "CH_LOG_LEVEL": "INFO", "CH_TRACE_135": "1" if trace else "0",
    })
    if localappdata is None:
        env.pop("LOCALAPPDATA", None)
    else:
        env["LOCALAPPDATA"] = str(tmp_path / localappdata) if localappdata else ""
    base = Path(env["LOCALAPPDATA"]) if env.get("LOCALAPPDATA") else home / "AppData" / "Local"
    env["CH_TEST_EXPECTED"] = str(base / "CHKanban" / "output")
    result = subprocess.run(
        [sys.executable, "-B", "-c", _CHILD_SETUP + textwrap.dedent(code)],
        cwd=tmp_path, env=env, capture_output=True, text=True, encoding="utf-8", timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    return result.stdout


def test_localappdata_is_user_specific(tmp_path):
    code = """
        from src.models.constants import LOCAL_OUTPUT_DIR
        assert isinstance(LOCAL_OUTPUT_DIR, str)
        assert Path(LOCAL_OUTPUT_DIR) == expected, LOCAL_OUTPUT_DIR
        print(json.dumps(LOCAL_OUTPUT_DIR))
    """
    paths = [json.loads(_run_child(tmp_path, code, localappdata=user)) for user in ("user-A", "user-B")]
    assert paths[0] != paths[1]


@pytest.mark.parametrize("localappdata", [None, ""], ids=["unset", "empty"])
def test_missing_localappdata_uses_current_home(tmp_path, localappdata):
    _run_child(tmp_path, """
        from src.models.constants import LOCAL_OUTPUT_DIR
        assert Path(LOCAL_OUTPUT_DIR) == Path.home() / 'AppData' / 'Local' / 'CHKanban' / 'output'
        assert Path(LOCAL_OUTPUT_DIR) == expected, LOCAL_OUTPUT_DIR
    """, localappdata=localappdata)


def test_local_output_is_not_developer_application_directory(tmp_path):
    _run_child(tmp_path, """
        from src.models import constants
        output = Path(constants.LOCAL_OUTPUT_DIR)
        assert output != Path(constants.__file__).resolve().parents[2]
        assert output.is_relative_to(sandbox), str(output)
    """)


def test_import_has_no_filesystem_side_effects(tmp_path):
    _run_child(tmp_path, """
        mutations = []
        def record(event, args):
            if event in ('os.mkdir', 'os.remove', 'os.rmdir', 'os.rename'):
                mutations.append(event)
            if event == 'open' and args[2] & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC):
                mutations.append(event)
        sys.addaudithook(record)
        from src.models.constants import LOCAL_OUTPUT_DIR
        assert not mutations, mutations
        assert not expected.parent.exists()
    """)


@pytest.mark.parametrize("trace", [False, True], ids=["normal", "trace"])
def test_logging_uses_local_output_and_creates_parents(tmp_path, trace):
    _run_child(tmp_path, """
        import logging
        from src.app import gui
        assert Path(gui.LOCAL_OUTPUT_DIR) == expected, gui.LOCAL_OUTPUT_DIR
        assert not expected.exists()
        gui._setup_logging()
        logging.getLogger('local-output-test').info('sandbox log marker')
        logging.shutdown()
        name = 'trace_135.log' if os.environ['CH_TRACE_135'] == '1' else 'ch_kanban.log'
        log = expected / 'logs' / name
        assert 'sandbox log marker' in log.read_text(encoding='utf-8')
        assert not (sandbox / 'logs').exists()
    """, trace=trace)


def test_logging_write_failure_does_not_fallback(tmp_path):
    _run_child(tmp_path, """
        from src.app import gui
        assert Path(gui.LOCAL_OUTPUT_DIR) == expected, gui.LOCAL_OUTPUT_DIR
        with patch.object(Path, 'mkdir', side_effect=PermissionError('denied')) as mkdir:
            try:
                gui._setup_logging()
            except PermissionError:
                pass
            else:
                raise AssertionError('write failure must propagate')
            assert mkdir.call_count == 1
        assert not expected.exists()
        assert not (sandbox / 'logs').exists()
    """)


@pytest.mark.parametrize("explicit_archive", [False, True], ids=["default-archive", "explicit-archive"])
def test_gui_saves_local_files_and_keeps_spo_export_dir(tmp_path, explicit_archive):
    _run_child(tmp_path, f"explicit_archive = {explicit_archive!r}\n" + textwrap.dedent("""
        from types import SimpleNamespace
        import pandas as pd
        from openpyxl import load_workbook
        from src.app import gui
        from src.services import exporter

        assert Path(gui.LOCAL_OUTPUT_DIR) == expected, gui.LOCAL_OUTPUT_DIR
        watch = sandbox / 'export'
        configured = sandbox / 'explicit-archive' if explicit_archive else None
        row = {
            '山通番': 1, 'ストア': 'TEST', 'OData_納入先': '元町', '納入先': '元町',
            'NONYUHIBIN': '2026060102', 'UKEIRE': '9P', 'SEBANGO': 'TEST', '番号': 1,
        }
        grouped = json.dumps([row], ensure_ascii=False)
        spo = pd.DataFrame([{
            'タイトル': '山1', '工程': '1工程', 'GroupedData': grouped, 'groupdata': grouped,
            'グループ番号': 1, 'パレット数': 1, 'Max移動工数': 0, '引取工数': 0,
            '引取開始時間': '',
        }])
        master = pd.DataFrame([{'OData_納入先': '元町-9P', 'NONYUHIBIN': '01', '入車時間': '07:30'}])
        app = SimpleNamespace(
            export_dir=str(watch), archive_enabled=True, archive_dir=str(configured) if configured else '',
            all_mountain_details_display=pd.DataFrame([row]), all_mountain_details=pd.DataFrame([row]),
            proc_details=pd.DataFrame(), mountain_proc=pd.DataFrame(), mountain_proc_map={},
            mountain_start_times={}, _collect_master_from_tree=lambda: master,
        )
        def staged(frame, watch_dir):
            return exporter.export_spo_xlsx_staged(frame, watch_dir, staging_dir=str(sandbox / 'staging'))
        with patch.object(gui, 'build_spo_export_df', return_value=spo), \
             patch.object(gui, 'get_master_path', return_value=sandbox / 'missing-master.xlsx'), \
             patch.object(gui, 'load_pickup_time_master_xlsx', return_value=master), \
             patch.object(gui, 'load_config', return_value={}), \
             patch.object(gui, 'export_spo_xlsx_staged', side_effect=staged), \
             patch.object(gui.messagebox, 'showinfo'), patch.object(gui.messagebox, 'showwarning') as warning:
            gui.App._auto_export_spo(app)
            assert not warning.called, warning.call_args_list

        history = expected / 'SPOアップロード用_履歴.xlsx'
        unmatched = expected / 'SPOアップロード用_未ヒット一覧.csv'
        assert len(pd.read_excel(history)) == 1
        assert pd.read_csv(unmatched, encoding='utf-8-sig').iloc[0]['vendor'] == '元町-9P'
        files = list(watch.iterdir())
        assert len(files) == 1 and files[0].suffix == '.xlsx'
        assert files[0].name.startswith('SPOアップロード用_')
        book = load_workbook(files[0])
        try:
            assert 'SPOExport' in book.active.tables
        finally:
            book.close()
        archive_root = configured if configured else expected / '_export_archive'
        manifest_path, = archive_root.glob('*/manifest.json')
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        assert manifest['出力ファイル名'] == files[0].name
        assert (manifest_path.parent / files[0].name).read_bytes() == files[0].read_bytes()
        if configured:
            assert not (expected / '_export_archive').exists()
        assert not list(expected.glob('SPOアップロード用_20*.xlsx'))
    """))


def test_current_audit_tools_follow_local_output(tmp_path):
    _run_child(tmp_path, """
        from tools import audit_spo_history, measure_p1_bundle_key_impact, verify_bundle_key_split_reason
        assert audit_spo_history._default_history_path() == expected / 'SPOアップロード用_履歴.xlsx'
        assert Path(measure_p1_bundle_key_impact.LOCAL_OUTPUT_DIR) == expected
        assert Path(verify_bundle_key_split_reason.LOCAL_OUTPUT_DIR) == expected
    """)


def test_split_reason_tool_creates_output_parent(tmp_path, monkeypatch):
    from collections import Counter
    import pandas as pd
    from tools import verify_bundle_key_split_reason as tool

    output = tmp_path / 'new' / 'output'
    monkeypatch.setattr(tool, 'LOCAL_OUTPUT_DIR', str(output))
    monkeypatch.setattr(tool, '_rebuild_details', lambda: pd.DataFrame())
    monkeypatch.setattr(tool, 'audit', lambda _: (pd.DataFrame(columns=['パレット行数増分']), Counter(), []))
    tool.main()
    assert (output / 'bundle_key_split_reason_audit.csv').is_file()
