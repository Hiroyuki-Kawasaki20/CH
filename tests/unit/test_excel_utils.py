# -*- coding: utf-8 -*-
"""CHかんばんセット — Excelユーティリティのユニットテスト"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.utils.excel_utils import _add_table_exact


def test_add_table_exact_is_idempotent_for_same_table_name(tmp_path):
    out_path = tmp_path / "idempotent.xlsx"
    pd.DataFrame([{"A": 1, "B": 2}]).to_excel(out_path, index=False, engine="openpyxl")

    _add_table_exact(out_path, "SPOExport")
    _add_table_exact(out_path, "SPOExport")

    ws = load_workbook(out_path).active
    assert list(ws.tables.keys()) == ["SPOExport"]
